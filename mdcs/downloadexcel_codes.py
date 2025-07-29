import os
import numpy as np
import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
import xmlschema
import re
import xml.etree.ElementTree as ET
from lxml import etree
import pickle
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from django.conf import settings


def load_excel_template(filepath: str, sheet_name: str, header_levels: tuple, drop_column_condition) -> pd.DataFrame:
    df = pd.read_excel(filepath, sheet_name=sheet_name, header=header_levels)
    columns_to_drop = [col for col in df.columns if drop_column_condition(col)]
    df = df.drop(columns=columns_to_drop)
    df = df[df.iloc[:, 0].notna()]
    df = df.drop(index=1, errors='ignore')
    return df

def load_pickle_mapping(pickle_path: str, tests: list) -> dict:
    with open(pickle_path, "rb") as f:
        AME = pickle.load(f)
    AME_rev = {}
    itr = 0
    for k, v in AME.items():
        if v[0] == "Notes":
            AME_rev[f"{tests[itr]}_Notes"] = k
            itr += 1
        else:
            AME_rev[v[0]] = k
#    for key, value in AME_rev.items():
#        print(f"{key}: {value}\n")       
    return AME_rev

def extract_paths_and_values(element, current_path=""):
    indexed_tags = ["Additive", "Replication"]
    semicolon_tags = ["Co-Authors","Point", "AdditionalProperties", "OtherMixingProperty", "Case"]
    values = {}
    children_by_tag = defaultdict(list)
    for child in element:
        children_by_tag[child.tag].append(child)
    for tag, children in children_by_tag.items():
        apply_indexing = tag in indexed_tags
        apply_semicolon = tag in semicolon_tags
        if apply_semicolon:
            combined_values = defaultdict(list)
            for child in children:
                if list(child):
                    sub_values = extract_paths_and_values(child, "")
                    for k, v in sub_values.items():
                        combined_values[k].append(v)
                elif child.text and child.text.strip():
                    combined_values["text"].append(child.text.strip())
            for k, vlist in combined_values.items():
                key_path = f"{current_path}.{tag}.{k}" if k != "text" else f"{current_path}.{tag}"
                values[key_path] = "; ".join(vlist)
        else:
            for idx, child in enumerate(children, 1):
                if tag == "Results" and "Stiffness" in current_path:
                    indexed_tag = f"{tag}_{idx}"
                else:
                    indexed_tag = f"{tag}_{idx}" if apply_indexing and len(children) > 1 else tag
                path = f"{current_path}.{indexed_tag}" if current_path else indexed_tag
                if list(child):
                    values.update(extract_paths_and_values(child, path))
                elif child.text and child.text.strip():
                    if child.tag == 'Notes' and not current_path:
                        path = f"{element.tag}_Notes"
                    values[path] = child.text.strip()
    return values

def fill_excel_from_xml_tree(xml_content: str, df_template: pd.DataFrame, reversed_dict: dict) -> pd.DataFrame:
    root = ET.fromstring(xml_content)
    xml_values = extract_paths_and_values(root)
    df_template = df_template.astype('object')
    for path, value in xml_values.items():
        if path in reversed_dict:
            column = reversed_dict[path]
            if column in df_template.columns:
                df_template.at[0, column] = value
            else:
                print(f"Warning: Column {column} not found in Excel.")
        else:
            print(f"Info: Path '{path}' not in mapping, skipped.")
    cleaned_columns = pd.MultiIndex.from_tuples(
        tuple("" if str(level).startswith("Unnamed") else level for level in col)
        for col in df_template.columns
    )
    df_template.columns = cleaned_columns
    return df_template

def process_xml(xml_contents: list, df_template: pd.DataFrame, reversed_dict: dict) -> pd.DataFrame:
    all_cases_df = pd.DataFrame()
    for xml_content in xml_contents:
        filled_df = fill_excel_from_xml_tree(xml_content, df_template.copy(), reversed_dict)
        all_cases_df = pd.concat([all_cases_df, filled_df], ignore_index=True)
    return all_cases_df

def excel_column_to_index(col_str):
    exp = 0
    col_index = 0
    for char in reversed(col_str.upper()):
        col_index += (ord(char) - ord('A') + 1) * (26 ** exp)
        exp += 1
    return col_index

def index_to_excel_column(col_idx):
    col_str = ''
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        col_str = chr(65 + remainder) + col_str
    return col_str

def AM_Excel(df_filled: pd.DataFrame, predefined_excel_path: str, download_path: str, start_row: int = 12, start_col: int = 2, sheet_name: str = 'Database (Columns)'):
    wb = load_workbook(predefined_excel_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    if isinstance(df_filled.columns, pd.MultiIndex):
        df_filled.columns = ['_'.join([str(lvl) for lvl in col if lvl]) for col in df_filled.columns]
    offset_ranges = [
        ("A", "BI", 1),
        ("BI", "JU", 2),
        ("JU", "RK", 3),
        ("RK", "SK", 4),
        ("SK", "UY", 5),
        ("UY", "WG", 6),
        ("WG", "XM", 7),
        ("XM", "ZR", 8),
    ]
    ranges_idx = [(excel_column_to_index(start), excel_column_to_index(end), offset) for start, end, offset in offset_ranges]
    df_col_count = len(df_filled.columns)
    col_excel_positions = []
    col_cursor = start_col
    for col_idx in range(df_col_count):
        for start_idx, end_idx, offset in ranges_idx:
            if start_idx <= col_cursor < end_idx:
                position_in_range = col_cursor - start_idx
                excel_col = start_idx + (offset - 1) + position_in_range
                col_excel_positions.append(excel_col)
                break
        else:
            excel_col = col_cursor
            col_excel_positions.append(excel_col)
        col_cursor += 1
    for row_idx, (_, row) in enumerate(df_filled.iterrows()):
        for col_idx_zero, value in enumerate(row):
            cell= ws.cell(row=start_row + row_idx, column=col_excel_positions[col_idx_zero])
            cell.value = value
            cell.alignment = Alignment(horizontal='center', vertical='center') 
    wb.save(download_path)
    print(f"Copied DataFrame to Excel with corrected column offsets, saved to {download_path}")

def main(xml_contents):
    predefined_excel = os.path.join(settings.BASE_DIR, 'static', 'AsphaltMine Database Structure - Download.xlsx')
    sheet_name = 'Database (Columns)'
    header_levels = (0,1,2,3,4,5,6,7,8)
    pickle_path = os.path.join(settings.BASE_DIR, "mdcs", "AM_excel_mapping.pkl")
    tests = ["RuttingExp","MarshallExp","ITSExp","TSRSTExp","UTSTExp","StiffnessExp"]
    df_template = load_excel_template(predefined_excel, sheet_name, header_levels, drop_column_condition=lambda col: 'E' in col)
    print("Template loaded, shape:", df_template.shape)
    AME_rev = load_pickle_mapping(pickle_path, tests)
    filled_df = process_xml(xml_contents, df_template, AME_rev)
    filled_output_path = "Initial_Download.xlsx"
    filled_df.to_excel(filled_output_path, sheet_name="Data")
    print(f"Process completed, saved filled Excel to '{filled_output_path}'")
    Final_path = "Download_Final.xlsx"
    AM_Excel(df_filled=filled_df, predefined_excel_path=predefined_excel, download_path=Final_path)
